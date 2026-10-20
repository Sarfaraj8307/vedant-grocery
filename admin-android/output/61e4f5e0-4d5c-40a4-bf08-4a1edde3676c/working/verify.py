from docx import Document
import sys
d = Document(sys.argv[1])
h1 = [p.text for p in d.paragraphs if p.style.name == 'Heading 1']
h2 = [p.text for p in d.paragraphs if p.style.name == 'Heading 2']
print('Heading1 count:', len(h1))
for h in h1:
    print('  H1:', h)
print('Heading2 count:', len(h2))
for h in h2:
    print('  H2:', h)
print('total paragraphs:', len(d.paragraphs))
print('tables:', len(d.tables))
